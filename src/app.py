import logging
import os
import re
import threading
import json
import asyncio
import webbrowser
import time
from datetime import datetime
from pathlib import Path

import ddddocr
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from flask import Flask, render_template, jsonify, request

# 确定项目根目录（src 的上一级）
PROJECT_ROOT = Path(__file__).resolve().parent.parent

app = Flask(__name__, template_folder='web/templates', static_folder='web/static')

# 配置文件保存路径
CONFIG_FILE = PROJECT_ROOT / "excel" / "config.json"


# ------------------ 日志内存记录器 ------------------
class MemoryLogHandler(logging.Handler):
    """自定义日志处理器，将日志保存在内存列表中以供前端轮询"""
    def __init__(self):
        super().__init__()
        self.log_lock = threading.Lock()
        self.logs = []

    def emit(self, record):
        msg = self.format(record)
        level = record.levelname.lower()
        with self.log_lock:
            self.logs.append({
                "text": msg,
                "level": level
            })

    def get_new_logs(self, start_index):
        with self.log_lock:
            total = len(self.logs)
            if start_index < total:
                return self.logs[start_index:], total
            return [], total

    def clear(self):
        with self.log_lock:
            self.logs.clear()


memory_logger = MemoryLogHandler()


# ------------------ 全局状态管理器 ------------------
class ScraperState:
    """线程安全的全局进度状态存储"""
    def __init__(self):
        self.lock = threading.Lock()
        self.students = []
        self.is_running = False
        self.config = {}

    def update_student(self, row, status, scores=None, message=""):
        with self.lock:
            for s in self.students:
                if s["row"] == row:
                    s["status"] = status
                    if scores is not None:
                        s["scores"] = scores
                    if message:
                        s["message"] = message
                    break

    def update_all_pending_to_error(self, message):
        with self.lock:
            for s in self.students:
                if s["status"] == "pending":
                    s["status"] = "error"
                    s["message"] = message


global_state = ScraperState()


class GlobalScraperRef:
    """持有 Scraper 对象的全局引用，用于通过 API 触发停止信号"""
    def __init__(self):
        self.scraper = None


global_scraper_ref = GlobalScraperRef()


# ------------------ Excel 数据读取与持久化 ------------------
def load_saved_config():
    if CONFIG_FILE.exists():
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def save_config(config_data):
    CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(config_data, f, ensure_ascii=False, indent=4)
    except Exception:
        pass


def load_excel_students(path, name_col, exam_id_col, pwd_col, score_col):
    """只读模式快速解析 Excel 表格，返回考生预览数据"""
    try:
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        students = []
        # 从第二行（跳过表头）开始读到最大行
        for r in range(2, ws.max_row + 1):
            name_val = ws.cell(row=r, column=name_col).value
            exam_id_val = ws.cell(row=r, column=exam_id_col).value
            pwd_val = ws.cell(row=r, column=pwd_col).value

            # 如果全是空单元格，视为无效行
            if name_val is None and exam_id_val is None and pwd_val is None:
                continue

            name = str(name_val).strip() if name_val is not None else ""
            exam_id = str(exam_id_val).strip() if exam_id_val is not None else ""
            password = str(pwd_val).strip() if pwd_val is not None else ""

            # 载入已有分数
            scores = []
            col = score_col
            while True:
                score_val = ws.cell(row=r, column=col).value
                if score_val is None:
                    break
                scores.append(score_val)
                col += 1

            students.append({
                "row": r,
                "name": name,
                "exam_id": exam_id,
                "password": password,
                "status": "pending",
                "scores": scores,
                "message": ""
            })
        wb.close()
        return students
    except Exception as e:
        print("读取 Excel 预览失败: ", e)
        return []


# ------------------ 网页端控制路由 API ------------------
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/config', methods=['GET', 'POST'])
def handle_config():
    if request.method == 'POST':
        config_data = request.json
        save_config(config_data)
        return jsonify({"status": "success"})
    else:
        saved = load_saved_config()
        # 默认项配置
        defaults = {
            "excel_path": "./excel/报名号.xlsx",
            "target_url": "https://cx.shmeea.edu.cn/shmeea/q/hgk2025yswquery3zd8#",
            "column_mapping": {
                "name": 2,
                "exam_id": 3,
                "password": 5,
                "first_score": 6
            },
            "start_row": 2,
            "max_row": 271,
            "overwrite": False,
            "export_new": False,
            "export_path": "./excel/导出成绩.xlsx",
            "headless": False,
            "selectors": {
                "exam_id": 'input[name="BMH"], input[name="bmh"], input[name="ZKZH"], input[name="zkzh"], #BMH, #ZKZH',
                "password": 'input[type="password"], #MM, #mm, input[name="MM"], input[name="mm"]',
                "captcha_input": 'input[name="verifyCode"], input[name="verifycode"], input[name="yzm"], #verifyCode',
                "captcha_img": 'img#verify, #verify, img[src*="verify"], img[src*="code"]',
                "query_btn": 'a:has-text("查询"), button:has-text("查询"), input[type="submit"], #queryBtn'
            },
            "extraction_mode": "自动智能匹配",
            "extraction_param": "无参数"
        }
        # 将已存盘的配置覆盖默认参数
        for k, v in saved.items():
            if isinstance(v, dict) and k in defaults:
                defaults[k].update(v)
            else:
                defaults[k] = v
        return jsonify(defaults)


@app.route('/api/upload', methods=['POST'])
def handle_upload():
    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "未找到上传的文件"})
    file = request.files['file']
    if file.filename == '':
        return jsonify({"status": "error", "message": "文件名为空"})

    excel_dir = PROJECT_ROOT / "excel"
    excel_dir.mkdir(parents=True, exist_ok=True)
    save_path = excel_dir / file.filename
    try:
        file.save(save_path)
        return jsonify({"status": "success", "path": str(save_path)})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)})


@app.route('/api/excel/preview')
def handle_preview():
    path = request.args.get("path")
    if not path or not Path(path).exists():
        return jsonify({"status": "error", "message": "指定的数据表不存在"})

    try:
        name_col = int(request.args.get("name_col", 2))
        id_col = int(request.args.get("id_col", 3))
        pwd_col = int(request.args.get("pwd_col", 5))
        score_col = int(request.args.get("score_col", 6))
    except ValueError:
        return jsonify({"status": "error", "message": "行列映射参数不合法"})

    students = load_excel_students(path, name_col, id_col, pwd_col, score_col)
    return jsonify({"status": "success", "students": students})


@app.route('/api/status')
def handle_status():
    try:
        last_log = int(request.args.get("last_log", 0))
    except ValueError:
        last_log = 0

    new_logs, total_logs = memory_logger.get_new_logs(last_log)

    with global_state.lock:
        students_copy = list(global_state.students)
        is_running = global_state.is_running

    return jsonify({
        "is_running": is_running,
        "students": students_copy,
        "logs": new_logs,
        "last_index": total_logs
    })


@app.route('/api/start', methods=['POST'])
def start_scraper():
    if global_state.is_running:
        return jsonify({"status": "error", "message": "查分系统正在运行中，无法重复启动"})

    config = request.json
    save_config(config)

    global_state.config = config
    global_state.is_running = True
    memory_logger.clear()

    # 初始化状态
    excel_students = load_excel_students(
        config["excel_path"],
        config["column_mapping"]["name"],
        config["column_mapping"]["exam_id"],
        config["column_mapping"]["password"],
        config["column_mapping"]["first_score"]
    )

    start_row = config["start_row"]
    max_row = config["max_row"]

    filtered_students = []
    for s in excel_students:
        if start_row <= s["row"] <= max_row:
            s["status"] = "pending"
            filtered_students.append(s)

    global_state.students = filtered_students

    # 创建独立子线程运行查分爬虫
    def scraper_runner():
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            scraper = ScoreScraper(config, global_state)
            global_scraper_ref.scraper = scraper
            scraper.run()
        except Exception as e:
            logging.critical("后台爬虫发生不可恢复的致命异常: %s", str(e), exc_info=True)
        finally:
            global_state.is_running = False
            loop.close()

    t = threading.Thread(target=scraper_runner, daemon=True)
    t.start()

    return jsonify({"status": "started"})


@app.route('/api/stop', methods=['POST'])
def stop_scraper():
    if global_scraper_ref.scraper:
        global_scraper_ref.scraper.stop_requested = True
        return jsonify({"status": "success", "message": "停止请求已送达"})
    return jsonify({"status": "error", "message": "查分系统当前并未运行"})


# ------------------ 本地系统文件浏览 Hook ------------------
@app.route('/api/browse/excel')
def browse_excel():
    import subprocess
    import sys
    script = """
import tkinter as tk
from tkinter import filedialog
root = tk.Tk()
root.withdraw()
root.attributes('-topmost', True)
path = filedialog.askopenfilename(filetypes=[('Excel文件', '*.xlsx')])
print(path)
"""
    try:
        # 使用 CREATE_NO_WINDOW 隐藏控制台窗口 (Windows specific: 0x08000000)
        creationflags = 0x08000000 if sys.platform == 'win32' else 0
        result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, creationflags=creationflags)
        path = result.stdout.strip()
        return jsonify({"path": path})
    except Exception as e:
        return jsonify({"path": "", "error": str(e)})


@app.route('/api/browse/folder')
def browse_folder():
    import subprocess
    import sys
    script = """
import tkinter as tk
from tkinter import filedialog
root = tk.Tk()
root.withdraw()
root.attributes('-topmost', True)
path = filedialog.asksaveasfilename(defaultextension='.xlsx', filetypes=[('Excel文件', '*.xlsx')])
print(path)
"""
    try:
        creationflags = 0x08000000 if sys.platform == 'win32' else 0
        result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, creationflags=creationflags)
        path = result.stdout.strip()
        return jsonify({"path": path})
    except Exception as e:
        return jsonify({"path": "", "error": str(e)})


# ------------------ 查分核心逻辑重构 (ScoreScraper) ------------------
class ScoreScraper:

    def __init__(self, user_config, state_manager):
        self.user_config = user_config
        self.config = self.get_merged_config()
        self.ocr = ddddocr.DdddOcr(show_ad=False)
        self.stop_requested = False
        self.state_manager = state_manager
        self.workbook = None

        self._setup_directories()
        self._setup_logging()
        self.logger = logging.getLogger(__name__)

    def get_merged_config(self):
        default_dirs = {
            "screenshot_dir": PROJECT_ROOT / "verify",
            "html_temp_dir": PROJECT_ROOT / "html",
            "selected_html_dir": PROJECT_ROOT / "selected_html",
            "scores_dir": PROJECT_ROOT / "scores",
            "log_dir": PROJECT_ROOT / "logs",
        }
        return {**default_dirs, **self.user_config}

    def _setup_directories(self):
        for dir_key in ["screenshot_dir", "html_temp_dir",
                        "selected_html_dir", "scores_dir", "log_dir"]:
            self.config[dir_key].mkdir(parents=True, exist_ok=True)

    def _setup_logging(self):
        log_file = self.config["log_dir"] / f"system_{datetime.now().strftime('%Y%m%d_%H%M')}.log"
        
        root_logger = logging.getLogger()
        root_logger.setLevel(logging.INFO)
        for handler in root_logger.handlers[:]:
            root_logger.removeHandler(handler)

        formatter = logging.Formatter("%(asctime)s - %(levelname)s - %(message)s", "%Y-%m-%d %H:%M:%S")

        fh = logging.FileHandler(log_file, encoding="utf-8")
        fh.setFormatter(formatter)
        
        sh = logging.StreamHandler()
        sh.setFormatter(formatter)
        
        memory_logger.setFormatter(formatter)

        root_logger.addHandler(fh)
        root_logger.addHandler(sh)
        root_logger.addHandler(memory_logger)

    def _extract_scores_from_html(self, html_content: str, exam_id: str, password: str, page_obj=None) -> list:
        mode = self.config["extraction_mode"]
        param = self.config["extraction_param"]

        self.logger.info("采用提取模式: %s (参数: %s)", mode, param)

        if mode == "行号切片范围":
            try:
                start, end = map(int, param.split(':'))
            except Exception:
                start, end = 62, 78
                self.logger.warning("行号切片参数格式不正确 (应如 62:78)，采用默认值 62:78")
            
            lines = html_content.splitlines()
            if start >= len(lines):
                self.logger.error("行号切片起始位置超出网页行数限制")
                return []
            selected_lines = lines[start:min(end, len(lines))]
            
            scores = []
            for line in selected_lines:
                matches = re.findall(r'>\s*(\d+(?:\.\d+)?)\s*<', line)
                for val_str in matches:
                    val = float(val_str)
                    if val.is_integer():
                        scores.append(int(val))
                    else:
                        scores.append(val)
            return scores

        elif mode == "自定义 CSS 选择器":
            if not page_obj:
                self.logger.error("CSS 模式下需要 Playwright Page 对象")
                return []
            try:
                elements = page_obj.locator(param).all()
                scores = []
                for el in elements:
                    text = el.inner_text().strip()
                    match = re.search(r'^\d+(?:\.\d+)?$', text)
                    if match:
                        val = float(text)
                        if val.is_integer():
                            scores.append(int(val))
                        else:
                            scores.append(val)
                return scores
            except Exception as e:
                self.logger.error("通过 CSS 选择器 %s 提取分数失败: %s", param, str(e))
                return []

        else:  # "自动智能匹配"
            # 过滤排除年份、考生账号、密码等，智能保留 0-760 之间的纯分数值
            raw_matches = re.findall(r'>\s*(\d+(?:\.\d+)?)\s*<', html_content)
            
            scores = []
            for val_str in raw_matches:
                if exam_id and val_str == exam_id:
                    continue
                if password and val_str == password:
                    continue

                val = float(val_str)
                # 排除 2000 到 2030 的年份数值
                if 2000 <= val <= 2030:
                    continue

                # 排除异常超大字符和不合理数字 (总分一般 <= 760)
                if 0 <= val <= 760 and len(val_str) <= 5:
                    if '.' in val_str:
                        if val.is_integer():
                            scores.append(int(val))
                        else:
                            scores.append(val)
                    else:
                        scores.append(int(val_str))
            return scores

    def _process_student(self, browser, row: int):
        ws = self.workbook.active
        exam_id_val = ws.cell(row=row, column=self.config["column_mapping"]["exam_id"]).value
        password_val = ws.cell(row=row, column=self.config["column_mapping"]["password"]).value
        name_val = ws.cell(row=row, column=self.config["column_mapping"]["name"]).value

        exam_id = str(exam_id_val).strip() if exam_id_val is not None else ""
        password = str(password_val).strip() if password_val is not None else ""
        name = str(name_val).strip() if name_val is not None else f"Row_{row}"

        if not exam_id or not password:
            self.logger.warning("跳过第 %d 行 (%s)：准考证号/报名号或密码为空", row, name)
            self.state_manager.update_student(row, "skipped", message="准考证号/密码为空")
            return

        first_score_col = self.config["column_mapping"]["first_score"]
        has_score = ws.cell(row=row, column=first_score_col).value is not None

        if has_score and not self.config["overwrite"]:
            self.logger.info("跳过考生 %s：成绩已存在，设置不覆盖", name)
            self.state_manager.update_student(row, "skipped", message="已有成绩不覆盖")
            return

        self.logger.info("========================================")
        self.logger.info("正在查询考生: %s (%s)", name, exam_id)
        self.state_manager.update_student(row, "running")

        context = None
        page = None
        try:
            # 只在学生查询这一级开启 Context/Page，保持单浏览器实例复用，极其轻量
            context = browser.new_context()
            page = context.new_page()
            page.set_default_timeout(15000)

            page.goto(self.config["target_url"])
            page.wait_for_timeout(500)

            # 输入准考证号
            page.locator(self.config["selectors"]["exam_id"]).first.click()
            page.locator(self.config["selectors"]["exam_id"]).first.fill(exam_id)
            page.wait_for_timeout(300)

            # 输入密码
            page.locator(self.config["selectors"]["password"]).first.click()
            page.locator(self.config["selectors"]["password"]).first.fill(password)
            page.wait_for_timeout(300)

            # 验证码重试环
            success = False
            for attempt in range(1, 4):
                if self.stop_requested:
                    self.logger.info("检测到停止信号，停止当前重试循环")
                    break

                self.logger.info("正在尝试识别验证码并登录 (第 %d 次)...", attempt)
                verify_element = page.locator(self.config["selectors"]["captcha_img"]).first
                
                if not verify_element.is_visible():
                    self.logger.warning("未找到验证码图片元素，尝试直接提交查询...")
                    page.locator(self.config["selectors"]["query_btn"]).first.click()
                    page.wait_for_timeout(1500)
                    success = True
                    break

                # 验证码截图与识别
                verify_png = self.config["screenshot_dir"] / "verify.png"
                verify_element.screenshot(path=verify_png)

                with open(verify_png, "rb") as f:
                    verify_code = self.ocr.classification(f.read())
                
                self.logger.info("验证码识别结果: %s", verify_code)

                # 输入验证码并点击查询
                captcha_input = page.locator(self.config["selectors"]["captcha_input"]).first
                captcha_input.click()
                captcha_input.fill(verify_code)
                page.wait_for_timeout(300)

                page.locator(self.config["selectors"]["query_btn"]).first.click()
                page.wait_for_timeout(1500)  # 等待页面载入

                # 验证码错误页面判断与捕获
                if captcha_input.is_visible():
                    page_text = page.content()
                    if "验证码" in page_text and ("错" in page_text or "不对" in page_text or "无效" in page_text):
                        self.logger.warning("验证码匹配失败，点击刷新重新尝试...")
                        verify_element.click()
                        page.wait_for_timeout(500)
                        continue
                    else:
                        self.logger.error("网页报错或密码账号错误，留在查询页无法进入。")
                        self.state_manager.update_student(row, "error", message="账号密码或网页错误")
                        break
                else:
                    success = True
                    break

            if success:
                # 截图保存成绩页面
                student_dir = self.config["scores_dir"] / f"{exam_id}_{name}"
                student_dir.mkdir(exist_ok=True)
                page.screenshot(path=student_dir / "full_page.png")
                self.logger.info("考生 %s 分数截图已保存", name)

                # 分数文本匹配提取
                html_content = page.content()
                scores = self._extract_scores_from_html(html_content, exam_id, password, page)

                if scores:
                    self.logger.info("考生 %s 提取分数: %s", name, str(scores))
                    for idx, score in enumerate(scores):
                        ws.cell(
                            row=row,
                            column=first_score_col + idx,
                            value=score
                        )
                    self.logger.info("考生 %s 成绩已存入 Excel", name)
                    self.state_manager.update_student(row, "success", scores=scores)
                else:
                    self.logger.warning("考生 %s 截图保存成功，但未能提取到任何分数数字", name)
                    self.state_manager.update_student(row, "success", message="仅保存截图")

        except Exception as e:
            self.logger.error("处理考生 %s 时遇到异常: %s", name, str(e), exc_info=True)
            self.state_manager.update_student(row, "error", message=str(e))
        finally:
            if page:
                try:
                    page.close()
                except Exception:
                    pass
            if context:
                try:
                    context.close()
                except Exception:
                    pass

    def run(self):
        self.logger.info("======== 系统开始运行 ========")
        try:
            self.workbook = load_workbook(self.config["excel_path"])
        except Exception as e:
            self.logger.critical("无法成功加载 Excel 工作簿: %s", str(e))
            self.state_manager.update_all_pending_to_error(str(e))
            return

        current_row = self.config["start_row"]
        max_row = self.config["max_row"]

        self.logger.info("正在初始化并启动 Playwright 引擎...")
        
        with sync_playwright() as playwright:
            browser = None
            try:
                browser = playwright.chromium.launch(headless=self.config["headless"])
                self.logger.info("Chromium 浏览器启动成功")
            except Exception as e:
                self.logger.critical("浏览器引擎启动失败，流程被中断: %s", str(e), exc_info=True)
                self.state_manager.update_all_pending_to_error(f"浏览器启动失败: {str(e)}")
                return

            try:
                while current_row <= max_row:
                    if self.stop_requested:
                        self.logger.info("收到用户停止指令，提前结束整个查询流程")
                        break

                    try:
                        self._process_student(browser, current_row)

                        if self.config["export_new"]:
                            save_path = Path(self.config["export_path"])
                            self.workbook.save(save_path)
                        else:
                            self.workbook.save(self.config["excel_path"])
                    except Exception as e:
                        self.logger.error("第 %d 行学生处理异常: %s", current_row, str(e))
                    finally:
                        current_row += 1
            finally:
                if browser:
                    browser.close()
                if self.workbook:
                    self.workbook.close()
                self.logger.info("======== 运行结束，资源已释放 ========")


# ------------------ Web 启动入口 ------------------
def open_browser():
    time.sleep(1.5)
    webbrowser.open("http://127.0.0.1:5000")


if __name__ == "__main__":
    # 在后台线程中自动打开浏览器
    threading.Thread(target=open_browser, daemon=True).start()
    # 启动 Flask 本地 Web 服务
    app.run(host="127.0.0.1", port=5000, debug=False)
